"""Load raw documents from the configured source directories as page texts."""
import hashlib
import logging
import os
import re
from dataclasses import dataclass, field

import pdfplumber

from .config import long_path

logging.getLogger("pdfminer").setLevel(logging.ERROR)


@dataclass
class Document:
    doc_id: str      # sha256 prefix of the file bytes; stable and filename-independent
    source: str
    filename: str    # original name; only ever written under private/
    pages: list
    kind: str = "pdf"
    ocr_pages: dict = field(default_factory=dict)   # 1-based page number -> OCR mean confidence

    @property
    def textless_pages(self):
        """1-based numbers of pages with no usable text (scanned images; not de-identified or published)."""
        if self.kind != "pdf":
            return []
        return [i for i, p in enumerate(self.pages, 1) if len(p.strip()) < TEXTLESS]


def _clean(text):
    text = text.replace(" ", " ").replace("\r", "")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _pdf_pages(path):
    with pdfplumber.open(path) as pdf:
        return [_clean(p.extract_text() or "") for p in pdf.pages]


def _txt_pages(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        return [_clean(p) for p in f.read().split("\f")]


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_text(el):
    out = []
    for node in el.iter():
        if node.tag == W + "t":
            out.append(node.text or "")
        elif node.tag == W + "tab":
            out.append("\t")
        elif node.tag in (W + "br", W + "cr"):
            out.append("\n")
    return "".join(out)


def _docx_pages(path):
    """Word via the standard library (a .docx is zipped XML). Tables become 'a | b | c' rows.
    Explicit page breaks split pages; otherwise the document is one page."""
    import xml.etree.ElementTree as ET
    import zipfile
    with zipfile.ZipFile(path) as z:
        body = ET.fromstring(z.read("word/document.xml")).find(W + "body")
        # Headers and footers can carry the document's date ("Last updated ..."): their text goes after the body.
        margins = []
        for name in sorted(n for n in z.namelist() if re.match(r"word/(header|footer)\d*\.xml$", n)):
            for p in ET.fromstring(z.read(name)).iter(W + "p"):
                t = _docx_text(p).strip()
                if t and t not in margins:
                    margins.append(t)
    pages, lines = [], []
    for block in body:
        if block.tag == W + "p":
            if any(br.get(W + "type") == "page" for br in block.iter(W + "br")) and lines:
                pages.append(lines)
                lines = []
            lines.append(_docx_text(block))
        elif block.tag == W + "tbl":
            for row in block.iter(W + "tr"):
                cells = [" ".join(_docx_text(p) for p in tc.iter(W + "p")).strip() for tc in row.iter(W + "tc")]
                lines.append(" | ".join(cells))
    pages.append(lines + margins)
    return [_clean("\n".join(p)) for p in pages]


def _cell(v):
    import datetime
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time() else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v).strip()


def _xlsx_pages(path):
    """One page per sheet; each row as 'a | b | c'. Trailing empty cells are dropped
    (some sheets report a used range 16k columns wide)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pages = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                rows.append(" | ".join(cells))
        pages.append(_clean(f"Sheet: {ws.title}\n" + "\n".join(rows)))
    wb.close()
    return pages


def _xml_pages(path):
    from . import ccda
    return ccda.read(path)   # None if not C-CDA -> counted as unsupported


READERS = {".pdf": _pdf_pages, ".txt": _txt_pages, ".docx": _docx_pages, ".xlsx": _xlsx_pages, ".xml": _xml_pages}
TEXTLESS = 50   # a page with fewer characters is treated as having no text layer (e.g. a scan)


TEXT_CACHE_VERSION = 1   # bump when a reader's output changes, to invalidate cached extractions


def ext_of(name):
    return os.path.splitext(name)[1].lower().lstrip(".")


def _cached_pages(reader, full, sha, ext, cache_dir):
    """pdfplumber takes ~0.1 s/page (~2 min for the corpus); extracted text is cached by file hash."""
    if cache_dir is None or ext != "pdf":
        return reader(full)
    import json
    f = cache_dir / f"{sha[:24]}-v{TEXT_CACHE_VERSION}.json"
    if f.exists():
        return json.loads(f.read_text(encoding="utf-8"))
    pages = reader(full)
    cache_dir.mkdir(parents=True, exist_ok=True)
    f.write_text(json.dumps(pages), encoding="utf-8")
    return pages


def load_documents(root, sources, ocr=None, text_cache=None):
    """ocr: an ocr.OCR instance; used only for sources with `ocr = true`."""
    docs, seen = [], set()
    stats = {"unsupported": 0, "duplicates": 0}
    for src in sources:
        directory = long_path(root / src["path"])
        for name in sorted(os.listdir(directory)):
            full = os.path.join(directory, name)
            reader = READERS.get(os.path.splitext(name)[1].lower())
            if reader is None or not os.path.isfile(full):
                stats["unsupported"] += 1
                continue
            with open(full, "rb") as f:
                raw = f.read()
            sha = hashlib.sha256(raw).hexdigest()
            if sha in seen:
                stats["duplicates"] += 1
                continue
            seen.add(sha)
            pages = _cached_pages(reader, full, sha, ext_of(name), text_cache)
            if pages is None:
                stats["unsupported"] += 1
                continue
            ext = os.path.splitext(name)[1].lower().lstrip(".")
            doc = Document(sha[:10], src["name"], name, pages, ext)
            if src.get("ocr") and ext == "pdf":
                if ocr is None:
                    raise RuntimeError(f"source '{src['name']}' needs OCR but OCR is not configured")
                for i, text in enumerate(pages):
                    if len(text.strip()) < TEXTLESS:
                        ocr_text, conf = ocr.page(raw, sha, i)
                        doc.pages[i] = _clean(ocr_text)
                        doc.ocr_pages[i + 1] = conf
            docs.append(doc)
    return docs, stats
