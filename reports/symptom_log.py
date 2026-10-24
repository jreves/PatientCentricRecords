"""The family's symptom log for the cycle in progress: a printable page (PDF) and the same sheet as a spreadsheet (.xlsx),
one row per day of the cycle, dates and dose days filled in from the regimen's schedule (cycles.py).

The visit notes record side effects only on visit days, so the days between doses, when counts are lowest, are mostly
blank; this sheet fills them in. A filled-in spreadsheet put with the records is read by the pipeline like any other
document, and its rows can then be cited for the case's ``symptom_log`` entries (shown as family reports).

Columns and the "when to call" line come from the case (``symptom_log_sheet``, see docs/case-file.md)."""
import html
from datetime import timedelta

from . import common

e = html.escape

COLUMNS = ["Temp (°F)", "Chills", "Fatigue (0–3)", "Nausea / vomiting", "Loss of appetite", "Bowel changes", "Bleeding / bruising"]
NOTES = "Other symptoms, extra medicines, notes"
CALL_IF = ("Call the care team as they have told you to: for example for a fever (often 100.4 °F / 38 °C or higher), chills, "
           "bleeding that won't stop, or vomiting or diarrhea that won't settle.")
WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def settings():
    s = common.case().get("symptom_log_sheet", {})
    cols = list(s.get("columns", COLUMNS)) + list(s.get("extra_columns", []))
    return cols, s.get("call_if", CALL_IF)


def rows(reg, c):
    """[(day, iso date, weekday, dose label)] for every day of cycle c."""
    out = []
    for day in range(1, reg.length + 1):
        d = c.start + timedelta(days=day - 1)
        out.append((day, d.isoformat(), WEEKDAYS[d.weekday()], f"day {day} dose" if day in reg.days else ""))
    return out


def title(reg, c):
    return f"Symptom log: {reg.drug}, cycle {c.n} ({common.fmt(c.start.isoformat())} – {common.fmt(c.last.isoformat())})"


CSS = common.CSS + """
@page { size: Letter portrait; margin: 9mm 8mm 10mm; }
body, .viz-root { background: #fff; } .viz-root { padding: 0; font-size: 8.5pt; } .wrap { max-width: none; }
h1 { font-size: 13pt; margin: 0 0 1mm; } .sub { font-size: 8pt; color: var(--text-secondary); margin: 0 0 1.5mm; max-width: none; }
.call { border: 1px solid var(--series-2); border-radius: 5px; padding: 1.2mm 2.5mm; font-size: 8pt; margin: 0 0 2mm; }
table.log { width: 100%; border-collapse: collapse; table-layout: fixed; font-size: 7.6pt; background: #fff; }
table.log th { font-weight: 600; font-size: 7pt; color: var(--text-secondary); border-bottom: 1.2px solid var(--text-secondary);
  padding: .8mm 1mm; vertical-align: bottom; white-space: normal; overflow-wrap: anywhere; hyphens: auto; }
table.log td { border-bottom: 1px solid var(--grid); border-right: 1px solid var(--grid); height: 6.9mm; padding: 0 1mm; }
table.log td.d { white-space: nowrap; color: var(--text-secondary); } table.log tr.dose td { background: #f1f7fd; }
table.log tr.dose td.dose { color: var(--text-primary); font-weight: 650; } table.log tr.wk td { border-top: 1.2px solid var(--guide); }
"""


def page(reg, c):
    cols, call_if = settings()
    head = ("<th style='width:4%'>Day</th><th style='width:9%'>Date</th><th style='width:10%'>Dose: given, held or late?</th>"
            + "".join(f"<th>{e(h)}</th>" for h in cols) + f"<th style='width:19%'>{e(NOTES)}</th>")
    body = []
    for day, iso, wd, dose in rows(reg, c):
        cls = " ".join(x for x in ("dose" if dose else "", "wk" if day > 1 and day % 7 == 1 else "") if x)
        body.append(f'<tr class="{cls}"><td class="d">{day}</td><td class="d">{wd} {int(iso[5:7])}/{int(iso[8:])}</td>'
                    f'<td class="d dose">{e(dose)}</td>' + "<td></td>" * (len(cols) + 1) + "</tr>")
    return (f'<h1>{e(title(reg, c))}</h1>'
            f'<p class="sub">{e(common.case().summary)}. One row a day; leave a box blank if there was nothing to note. Fatigue: '
            '0 = none, 1 = tired but managing, 2 = resting much of the day, 3 = in bed most of the day. Bring it to the next '
            'visit. To add it to the reports, fill in the spreadsheet version and put it with the records.</p>'
            f'<div class="call">{e(call_if)}</div>'
            f'<table class="log"><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table>')


def write_xlsx(reg, c, path):
    """The same sheet as a spreadsheet; dates are written as text (YYYY-MM-DD), so they read back exactly."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    cols, call_if = settings()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Cycle {c.n}"
    ws.append([title(reg, c)])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([call_if])
    ws.append([])
    header = ["Day", "Date", "Weekday", "Dose: given, held or late?"] + cols + [NOTES]
    ws.append(header)
    for cell in ws[4]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="bottom")
    shade = PatternFill("solid", fgColor="F1F7FD")
    for day, iso, wd, dose in rows(reg, c):
        ws.append([day, iso, wd, dose] + [None] * (len(cols) + 1))
        if dose:
            for cell in ws[ws.max_row]:
                cell.fill = shade
    for i, w in enumerate([5, 11, 9, 16] + [12] * len(cols) + [40], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "E5"
    wb.save(path)


def build(reg, slug):
    """The log for the cycle in progress (none once a regimen has ended). Returns the files written."""
    c = reg.current()
    if c is None:
        return []
    name = f"symptom_log_{slug}_cycle{c.n}"
    doc = (f'<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>Symptom log, cycle {c.n}</title><style>{CSS}</style></head>'
           f'<body><div class="viz-root print"><div class="wrap">{page(reg, c)}</div></div></body></html>')
    pdf = common.write_and_print(name, doc, f"{name}_{common.case().built}.pdf")
    xlsx = common.out_dir() / f"{name}.xlsx"
    write_xlsx(reg, c, xlsx)
    return [pdf, xlsx]
